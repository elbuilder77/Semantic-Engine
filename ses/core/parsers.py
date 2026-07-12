import csv
import io
import logging
from typing import BinaryIO

# Conditional imports for different file types
try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    import docx
except ImportError:
    docx = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

logger = logging.getLogger(__name__)

def extract_text_content(file_obj: BinaryIO, filename: str) -> str:
    """
    Detecta el tipo de archivo y extrae su contenido de texto.
    Soporta: PDF, DOCX, CSV, XLSX, TXT.
    """
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    
    try:
        if ext == "pdf":
            return _extract_pdf(file_obj)
        elif ext == "docx":
            return _extract_docx(file_obj)
        elif ext in ["csv", "xlsx"]:
            return _extract_tabular(file_obj, ext)
        elif ext == "xls":
            raise ValueError("Formato XLS legacy no soportado; convierta el archivo a XLSX.")
        elif ext in ["txt", "md", "json", "py", "rs"]:
            return file_obj.read().decode("utf-8", errors="ignore")
        else:
            # Fallback a lectura de texto plano para extensiones desconocidas
            try:
                content = file_obj.read().decode("utf-8")
                return content
            except UnicodeDecodeError:
                logger.warning("No se pudo determinar el extractor para %s", filename)
                return ""
    except Exception as e:
        logger.error("Error extrayendo texto de %s: %s", filename, e)
        return ""

def _extract_pdf(file_obj: BinaryIO) -> str:
    if PdfReader is None:
        raise RuntimeError("pypdf no instalado")
    
    text = []
    reader = PdfReader(file_obj)
    for page_number, page in enumerate(reader.pages, start=1):
        page_text = page.extract_text()
        if page_text:
            text.append(f"[page {page_number}]\n{page_text}")
    return "\n\n".join(text)

def _extract_docx(file_obj: BinaryIO) -> str:
    if docx is None:
        raise RuntimeError("python-docx no instalado")
    
    doc = docx.Document(file_obj)
    return "\n".join([para.text for para in doc.paragraphs])

def _extract_tabular(file_obj: BinaryIO, ext: str) -> str:
    if ext == "csv":
        return _extract_csv(file_obj)

    if openpyxl is None:
        raise RuntimeError("openpyxl no instalado")

    workbook = openpyxl.load_workbook(file_obj, read_only=True, data_only=True)
    rows = []
    for sheet in workbook.worksheets:
        rows.append(f"# Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in row]
            if any(values):
                rows.append(" | ".join(values))
    workbook.close()
    return "\n".join(rows)


def _extract_csv(file_obj: BinaryIO) -> str:
    raw = file_obj.read()
    if isinstance(raw, str):
        text_stream = io.StringIO(raw)
    else:
        text_stream = io.StringIO(raw.decode("utf-8-sig", errors="replace"))

    rows = []
    for row in csv.reader(text_stream):
        if any(cell.strip() for cell in row):
            rows.append(" | ".join(cell.strip() for cell in row))
    return "\n".join(rows)
